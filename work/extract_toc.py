from docx import Document
import glob, os
os.chdir(r"C:\Users\orang\Desktop\网页制作")
for f in glob.glob("*.docx"):
    try:
        doc = Document(f)
        print("=" * 80)
        print("FILE:", f)
        print("Paras:", len(doc.paragraphs), "Tables:", len(doc.tables))
        n = 0
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                print(f"  [{p.style.name}] {t[:90]}")
                n += 1
                if n > 45:
                    break
    except Exception as e:
        print("ERR", f, e)
