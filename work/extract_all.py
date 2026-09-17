from docx import Document
import os, json
os.chdir(r"C:\Users\orang\Desktop\网页制作")
targets = ["2025年纽约市美容院行业调研报告 最终版.docx", "2025纽约市食品进口行业调研报告 (终稿).docx", "2.HAO YUN LAI WHOLESALE INC 个案调研报告 （中文）.docx", "Reachin Capital Inc. 个案调研报告（2025）终稿.docx"]
out = {}
for f in targets:
    doc = Document(f)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables = []
    for t in doc.tables:
        rows = [[c.text.strip()[:80] for c in r.cells] for r in t.rows]
        tables.append(rows)
    n = len(paras)
    out[f] = {"n_paras": n, "n_tables": len(tables), "tables": tables,
              "head": paras[:25], "tail": paras[-15:],
              "mid": paras[n//2-7:n//2+8]}
with open("work/full_extract.json", "w", encoding="utf-8") as fh:
    json.dump(out, fh, ensure_ascii=False, indent=1)
print("saved", {k: v["n_paras"] for k, v in out.items()})
