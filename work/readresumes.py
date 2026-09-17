from docx import Document
import os
for f in ["董事资料/QI_GONG_Resume_English.docx", "董事资料/Keejong_Chang_EnglishResume_updated.docx"]:
    p = os.path.join(r"C:\Users\orang\Desktop\网页制作", f)
    doc = Document(p)
    print("=" * 60)
    print("FILE:", f)
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            print("  [" + para.style.name + "] " + t[:200])
    for t in doc.tables:
        print("  [TABLE]")
        for r in t.rows:
            print("   | " + " | ".join(c.text.strip()[:80] for c in r.cells))
